# Генерация шаблона сметы (.xlsx, RU/KZ/EN) с формулами → static/docs/templates/{lang}/
import json, os, subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
T = json.loads(subprocess.check_output(['node', '-e', "process.stdout.write(JSON.stringify(require('./texts')))"], cwd=HERE))

DP, VI, GO = '2B005B', '6E3FA3', 'D6A86A'
F = 'Arial'
thin = Side(style='thin', color='CFC8DA')
B = Border(left=thin, right=thin, top=thin, bottom=thin)
INPUT = PatternFill('solid', fgColor='FFF6D5')
HEAD = PatternFill('solid', fgColor=DP)
CAT = PatternFill('solid', fgColor='EFEAF6')
TOT = PatternFill('solid', fgColor=GO)
MONEY = '#,##0'

for lang, L in T.items():
    b = L['budget']
    wb = Workbook(); ws = wb.active; ws.title = b['sheet']
    widths = [6, 52, 10, 9, 16, 17, 17, 18, 40]
    for i, w in enumerate(widths): ws.column_dimensions['ABCDEFGHI'[i]].width = w

    ws['A1'] = L['fund']; ws['A1'].font = Font(name=F, size=9, bold=True, color=VI)
    ws['A2'] = b['title']; ws['A2'].font = Font(name=F, size=16, bold=True, color=DP)
    r = 4
    for m in b['meta']:
        ws.cell(r, 1, m).font = Font(name=F, size=10, bold=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=9)
        for c in range(3, 10):
            ws.cell(r, c).fill = INPUT; ws.cell(r, c).border = B
        ws.cell(r, 3).font = Font(name=F, size=10)
        r += 1
    r += 1
    ws.cell(r, 1, b['legend']).font = Font(name=F, size=9, italic=True, color='7A7486')
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[r].height = 28
    r += 2

    hr = r
    for i, h in enumerate(b['head'], 1):
        c = ws.cell(r, i, h); c.font = Font(name=F, size=10, bold=True, color='FFFFFF'); c.fill = HEAD; c.border = B
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[r].height = 32
    ws.freeze_panes = ws.cell(r + 1, 1)
    r += 1

    subtotal_rows = []
    for k, cat in enumerate(b['cats'], 1):
        # строка статьи
        ws.cell(r, 1, k).font = Font(name=F, size=10, bold=True)
        ws.cell(r, 2, cat).font = Font(name=F, size=10, bold=True, color=DP)
        for c in range(1, 10):
            ws.cell(r, c).fill = CAT; ws.cell(r, c).border = B
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='center')
        cat_row = r; r += 1
        first = r
        for j in range(3):  # 3 строки для позиций
            ws.cell(r, 1, f'{k}.{j + 1}').font = Font(name=F, size=10, color='7A7486')
            if k == 1 and j == 0:
                ex = b['ex']
                ws.cell(r, 2, ex[0]); ws.cell(r, 3, ex[1]); ws.cell(r, 4, ex[2]); ws.cell(r, 5, ex[3]); ws.cell(r, 9, ex[4])
                ws.cell(r, 7, f'=F{r}')
                for c in (2, 3, 4, 5, 7, 9): ws.cell(r, c).font = Font(name=F, size=10, italic=True, color='7A7486')
            for c in (2, 3, 4, 5, 7, 9):
                ws.cell(r, c).fill = INPUT
                if ws.cell(r, c).font.name != F or not ws.cell(r, c).font.italic:
                    ws.cell(r, c).font = Font(name=F, size=10, italic=(k == 1 and j == 0), color='7A7486' if (k == 1 and j == 0) else '000000')
            ws.cell(r, 6, f'=IF(AND(D{r}<>"",E{r}<>""),D{r}*E{r},0)')
            ws.cell(r, 8, f'=F{r}-G{r}')
            for c in range(1, 10):
                ws.cell(r, c).border = B
                ws.cell(r, c).alignment = Alignment(wrap_text=c in (2, 9), vertical='center')
            for c in (5, 6, 7, 8): ws.cell(r, c).number_format = MONEY
            ws.cell(r, 6).font = Font(name=F, size=10); ws.cell(r, 8).font = Font(name=F, size=10)
            r += 1
        last = r - 1
        for c in (6, 7, 8):
            col = 'ABCDEFGHI'[c - 1]
            ws.cell(cat_row, c, f'=SUM({col}{first}:{col}{last})').number_format = MONEY
            ws.cell(cat_row, c).font = Font(name=F, size=10, bold=True)
        subtotal_rows.append(cat_row)

    # итог
    ws.cell(r, 2, b['total']).font = Font(name=F, size=11, bold=True, color=DP)
    for c in range(1, 10):
        ws.cell(r, c).fill = TOT; ws.cell(r, c).border = B
    for c in (6, 7, 8):
        col = 'ABCDEFGHI'[c - 1]
        ws.cell(r, c, '=' + '+'.join(f'{col}{x}' for x in subtotal_rows)).number_format = MONEY
        ws.cell(r, c).font = Font(name=F, size=11, bold=True, color=DP)
    tot = r; r += 1
    ws.cell(r, 2, b['share']).font = Font(name=F, size=10, bold=True)
    ws.cell(r, 7, f'=IF(F{tot}>0,G{tot}/F{tot},0)').number_format = '0%'
    ws.cell(r, 7).font = Font(name=F, size=10, bold=True)

    # подписи
    r += 3
    for i, s in enumerate(L['sign'][:3]):
        ws.cell(r, 2 + i * 3 if i else 2, '_' * 24 if i == 0 else '_' * 18).font = Font(name=F, size=10)
        ws.cell(r + 1, 2 + i * 3 if i else 2, s).font = Font(name=F, size=8, color='7A7486')

    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f'{hr}:{hr}'

    out = os.path.join(HERE, '..', 'static', 'docs', 'templates', lang)
    os.makedirs(out, exist_ok=True)
    wb.save(os.path.join(out, b['file']))
    print(lang, 'xlsx ok')
